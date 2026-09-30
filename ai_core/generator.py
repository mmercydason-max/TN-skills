"""Formatting modules: raw AI text -> styled HTML preview, .docx and .pdf bytes."""
import html
import io
import re
from typing import Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from fpdf import FPDF
from fpdf.enums import XPos, YPos
from PIL import Image

import config
from ai_core.text_utils import (
    BULLET, HEADING, PARA, SUB, TITLE,
    parse_blocks, sanitize_text, split_terms,
)

__all__ = ["sanitize_text", "format_docx", "format_pdf", "format_html_preview"]

FONT_NAME = "Times New Roman"


def _split_title(text: str, doc_type: str):
    """Return (title, remaining_blocks)."""
    blocks = parse_blocks(text)
    if blocks and blocks[0][0] == TITLE:
        return blocks[0][1], blocks[1:]
    return (doc_type or "Legal Document").strip().title(), blocks


# --------------------------------------------------------------------------- HTML
def format_html_preview(text: str) -> str:
    """Convert the document to styled, HTML-escaped blocks (no blank lines/indents,
    so Streamlit's markdown parser leaves it alone)."""
    def esc(s: str) -> str:
        return html.escape(s).replace("$", "&#36;")   # '$' would trigger LaTeX in markdown

    parts = []
    for kind, content in parse_blocks(text):
        c = esc(content)
        if kind == TITLE:
            parts.append(f"<h2 style='text-align:center;margin:0 0 14px 0;'>{c}</h2>")
        elif kind == HEADING:
            parts.append(f"<h4 style='margin:18px 0 6px 0;color:#8ab4ff;'>{c}</h4>")
        elif kind == SUB:
            parts.append(f"<p style='margin:2px 0 2px 28px;line-height:1.55;'>{c}</p>")
        elif kind == BULLET:
            parts.append(f"<p style='margin:2px 0 2px 22px;line-height:1.55;'>&bull; {c}</p>")
        else:
            parts.append(f"<p style='margin:6px 0;line-height:1.6;text-align:justify;'>{c}</p>")
    return "".join(parts)


# --------------------------------------------------------------------------- DOCX
def _set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _add_page_number(paragraph):
    run = paragraph.add_run()
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(9)
    run.font.name = FONT_NAME


def format_docx(text: str, doc_type: str, terms: str = "") -> bytes:
    """Build a Word document: logo, title, optional Key Terms table, body, footer."""
    title, blocks = _split_title(sanitize_text(text), doc_type)
    doc = Document()

    section = doc.sections[0]
    section.left_margin = section.right_margin = Inches(1)
    section.top_margin = section.bottom_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)
    normal.paragraph_format.space_after = Pt(6)

    if config.LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(config.LOGO_PATH), width=Inches(2.4))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = FONT_NAME

    term_list = split_terms(terms)
    if term_list:
        h = doc.add_paragraph()
        r = h.add_run("Summary of Key Terms")
        r.bold = True
        r.font.size = Pt(13)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for cell, label in zip(hdr, ("No.", "Term / Condition")):
            cell.text = ""
            run = cell.paragraphs[0].add_run(label)
            run.bold = True
            _set_cell_shading(cell, "D9E2F3")
        for i, term in enumerate(term_list, 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = term
        table.autofit = False
        table.columns[0].width = Inches(0.6)
        table.columns[1].width = Inches(5.9)
        for row in table.rows:
            row.cells[0].width = Inches(0.6)
            row.cells[1].width = Inches(5.9)
        doc.add_paragraph()

    for kind, content in blocks:
        if kind == HEADING:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(content)
            r.bold = True
            r.font.size = Pt(13)
        elif kind == SUB:
            p = doc.add_paragraph(content)
            p.paragraph_format.left_indent = Inches(0.4)
        elif kind == BULLET:
            p = doc.add_paragraph(content, style="List Bullet")
        else:
            p = doc.add_paragraph(content)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(18)
    r = note.add_run(config.DISCLAIMER)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    doc.styles["Footer"].font.size = Pt(9)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run(config.FOOTER_TEXT + "  |  Page ")
    fr.font.size = Pt(9)
    fr.font.name = FONT_NAME
    fr.italic = True
    _add_page_number(footer)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------- PDF
_PDF_CHAR_MAP = {"\u20b9": "Rs. ", "\u20ac": "EUR ", "\u2122": "(TM)", "\u2192": "->"}


def _pdf_safe(text: str) -> str:
    """Core PDF fonts are Latin-1 only; map common symbols, replace the rest."""
    for src, dst in _PDF_CHAR_MAP.items():
        text = text.replace(src, dst)
    return text.encode("latin-1", "replace").decode("latin-1")


class _LegalPDF(FPDF):
    def __init__(self, title: str):
        super().__init__(orientation="P", unit="mm", format="A4")
        self.doc_title = _pdf_safe(title)
        self.set_margins(22, 20, 22)
        self.set_auto_page_break(auto=True, margin=22)
        self.alias_nb_pages()
        self.logo_w = 42.0
        self.logo_h = 0.0
        if config.LOGO_PATH.exists():
            with Image.open(config.LOGO_PATH) as im:
                self.logo_h = self.logo_w * im.height / im.width

    def header(self):
        y = 8.0
        if self.logo_h:
            self.image(str(config.LOGO_PATH), x=(self.w - self.logo_w) / 2, y=y, w=self.logo_w)
            y += self.logo_h + 1
        self.set_xy(self.l_margin, y)
        self.set_font("Helvetica", "B", 11)
        self.set_text_color(30, 30, 30)
        self.cell(0, 6, self.doc_title, align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.set_draw_color(180, 180, 180)
        self.line(self.l_margin, y + 7, self.w - self.r_margin, y + 7)
        self.set_y(y + 12)

    def footer(self):
        self.set_y(-16)
        self.set_font("Helvetica", "I", 8)
        self.set_text_color(110, 110, 110)
        self.cell(0, 5, _pdf_safe(config.FOOTER_TEXT), align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        self.cell(0, 5, f"Page {self.page_no()} of {{nb}}", align="C")


def format_pdf(text: str, doc_type: str, terms: str = "") -> bytes:
    """Branded PDF: centered logo + title on every page, bold headings, bullet terms, footer."""
    title, blocks = _split_title(sanitize_text(text), doc_type)
    pdf = _LegalPDF(title)
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin

    def para(txt, font_style="", size=11, indent=0.0, align="J", gap=2.0, line_h=5.8):
        pdf.set_font("Helvetica", font_style, size)
        pdf.set_text_color(20, 20, 20)
        pdf.set_x(pdf.l_margin + indent)
        pdf.multi_cell(width - indent, line_h, _pdf_safe(txt), align=align,
                       new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(gap)

    def bullet(txt, indent=4.0):
        pdf.set_font("Helvetica", "", 11)
        pdf.set_text_color(20, 20, 20)
        pdf.set_x(pdf.l_margin + indent)
        pdf.cell(5, 5.8, "-")
        pdf.multi_cell(width - indent - 5, 5.8, _pdf_safe(txt), align="L",
                       new_x=XPos.LMARGIN, new_y=YPos.NEXT)
        pdf.ln(1)

    term_list = split_terms(terms)
    if term_list:
        para("Summary of Key Terms", "B", 12, align="L", gap=1)
        for t in term_list:
            bullet(t)
        pdf.ln(4)

    for kind, content in blocks:
        if kind == HEADING:
            pdf.ln(3)
            para(content, "B", 12, align="L", gap=1)
        elif kind == SUB:
            para(content, indent=8)
        elif kind == BULLET:
            bullet(content)
        else:
            para(content)

    pdf.ln(4)
    para(config.DISCLAIMER, "I", 8, align="L")
    return bytes(pdf.output())
