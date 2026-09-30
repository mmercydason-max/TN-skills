"""Offline tests (no API key needed) for text helpers and DOCX/PDF/HTML exporters."""
import io

from docx import Document

from ai_core.generator import format_docx, format_html_preview, format_pdf
from ai_core.text_utils import parse_blocks, sanitize_text, split_terms

SAMPLE = """# Freelance Work Contract

## 1. Services
The \u201cService Provider\u201d shall deliver work for a fee of $5,000.
(a) design mockups
## 2. Term
- Either party may terminate with 15 days notice
IN WITNESS WHEREOF, the parties have signed.
Jane Doe (Service Provider)"""


def test_sanitize_removes_typographic_quotes_and_markdown():
    out = sanitize_text("\u201cHi\u201d \u2014 **bold** `x`")
    assert out == '"Hi" - bold x'


def test_split_terms_handles_semicolons_and_newlines():
    assert split_terms("A; B;\n- C") == ["A", "B", "C"]
    assert split_terms("") == []


def test_parse_blocks_kinds():
    kinds = [k for k, _ in parse_blocks(SAMPLE)]
    assert kinds[0] == "title" and "heading" in kinds and "sub" in kinds and "bullet" in kinds


def test_html_preview_escapes_and_has_no_blank_lines():
    html = format_html_preview("# T\n<script>alert(1)</script> costs $5")
    assert "<script>" not in html and "$" not in html and "\n" not in html


def test_docx_contains_title_terms_table_and_footer():
    data = format_docx(SAMPLE, "Freelance Work Contract", "Pay in 7 days; IP to client")
    doc = Document(io.BytesIO(data))
    assert any("Freelance Work Contract" in p.text for p in doc.paragraphs)
    assert len(doc.tables) == 1 and len(doc.tables[0].rows) == 3
    assert "All Rights Reserved" in doc.sections[0].footer.paragraphs[0].text


def test_docx_without_terms_has_no_table():
    doc = Document(io.BytesIO(format_docx(SAMPLE, "x", "")))
    assert len(doc.tables) == 0


def test_pdf_is_valid_and_multipage_for_long_text():
    long_text = SAMPLE + "\n" + "\n".join(f"Clause {i}. " + "Lorem ipsum dolor sit amet. " * 12 for i in range(60))
    data = format_pdf(long_text, "Freelance Work Contract", "Pay in 7 days; \u20b9 5,000 deposit")
    assert data.startswith(b"%PDF") and len(data) > 2000
